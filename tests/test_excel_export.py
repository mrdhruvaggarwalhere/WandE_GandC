"""
Unit tests for multi-sheet Excel export and Column A:G mapping.
"""

import io
import openpyxl
import pytest
from app.core.database import init_db
from app.core.seed_data import seed_all
from app.core.excel_exporter import generate_full_excel_workbook

def test_excel_export_structure():
    init_db()
    seed_all()

    excel_bytes = generate_full_excel_workbook()
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    # Verify all 8 mandatory sheets exist
    assert "Deals (A-G Mapping)" in sheet_names
    assert "Deal Chains (Lots)" in sheet_names
    assert "Billing Instructions" in sheet_names
    assert "Brokerage Detail" in sheet_names
    assert "Party Ledger & Dues" in sheet_names
    assert "Products" in sheet_names
    assert "Parties" in sheet_names
    assert "Audit Log" in sheet_names

    # Check Columns A:G on primary Deals sheet
    ws_deals = wb["Deals (A-G Mapping)"]
    headers = [ws_deals.cell(row=1, column=c).value for c in range(1, 8)]

    assert headers[0].startswith("Deal Date") # Col A
    assert headers[1].startswith("Buyer")     # Col B
    assert headers[2].startswith("Seller")    # Col C
    assert headers[3].startswith("Product")   # Col D
    assert headers[4].startswith("Quantity")  # Col E
    assert headers[5].startswith("Price and GST") # Col F
    assert headers[6].startswith("Delivery Date") # Col G

    # Check rows exist
    assert ws_deals.max_row >= 4 # Header + 3 seeded deals from worked example
